"""Full backup ZIP: right files, right numbers, no secrets, admins only."""
import io
import os
import tempfile
import time
import zipfile

os.environ["SF_NO_THREADS"] = "1"
_d = tempfile.mkdtemp()
os.environ["SF_DB"] = os.path.join(_d, "bk.db")
os.environ["SF_BACKUP"] = os.path.join(_d, "bk")

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import db, ops  # noqa: E402
from app.main import app  # noqa: E402

ADMIN = {"role": "admin", "id": 0, "name": "t", "language_id": None}


def test_full_backup_zip():
    with TestClient(app) as cl:
        with db.tx() as c:
            lid = c.execute("INSERT INTO languages(code,name,target,created_at) VALUES('BK','Backupese',0,1)").lastrowid
            a, acode = ops.create_user(c, ADMIN, "Asha", "participant", lid, phone="9100000031")
            b, _ = ops.create_user(c, ADMIN, "Bala", "participant", lid, phone="9100000032")
            sup, scode = ops.create_user(c, ADMIN, "Sup", "lead", lid, phone="9100000033")
            c.execute("UPDATE users SET password_hash='SECRETHASH', password_salt='SECRETSALT' WHERE id=?", (a,))
            pid = c.execute("INSERT INTO pairs(language_id,a,b,status,created_at) VALUES(?,?,?,?,?)", (lid, a, b, "ACTIVE", time.time())).lastrowid
            for i, (st, rv, mins) in enumerate((("ACCEPTED", "ACCEPTED", 12), ("REJECTED", "REJECTED", 6), ("SUBMITTED", "PENDING", 9), ("ASSIGNED", "PENDING", None))):
                sid = c.execute("INSERT INTO scripts(language_id,seq,code,domain,status,shingles,created_at) VALUES(?,?,?,?,?,?,?)",
                                (lid, 900 + i, f"BK-{i}", "BFSI", "READY", b"FINGERPRINTBLOB", time.time())).lastrowid
                c.execute("INSERT INTO assignments(script_id,version,pair_id,ua,ub,status,review_status,conf_min,submitted_at) VALUES(?,?,?,?,?,?,?,?,?)",
                          (sid, 1, pid, a, b, st, rv, mins, time.time() if mins else 0))
            admin_code = c.execute("SELECT code FROM users WHERE role='admin' ORDER BY id LIMIT 1").fetchone()["code"]
        assert cl.post("/login", data={"code": admin_code}).status_code == 200
        r = cl.get("/admin/backup/full.zip")
        assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
        z = zipfile.ZipFile(io.BytesIO(r.content))
        names = set(z.namelist())
        assert {"ScriptFlow_report.xlsx", "README.txt", "tables/users.csv", "tables/assignments.csv", "tables/scripts.csv",
                "tables/script_versions.csv", "tables/pairs.csv"} <= names
        assert not any(n.endswith(("api_keys.csv", "settings.csv", "sessions.csv", "connections.csv", "dev_keys.csv")) for n in names)
        everything = b"".join(z.read(n) for n in names if n.endswith(".csv"))
        for secret in (b"SECRETHASH", b"SECRETSALT", acode.encode(), admin_code.encode(), b"FINGERPRINTBLOB"):
            assert secret not in everything
        wb = openpyxl.load_workbook(io.BytesIO(z.read("ScriptFlow_report.xlsx")))
        assert wb.sheetnames == ["Summary", "Work done", "Agents", "Pairs", "Scripts", "Languages"]
        work = list(wb["Work done"].iter_rows(min_row=2, values_only=True))
        assert sorted(w[10] for w in work) == ["Accepted", "Pending", "Rejected"]   # the unsubmitted one is left out
        total = [row for row in wb["Summary"].iter_rows(values_only=True) if row[0] == "All languages"][0]
        assert total[1] == 3 and total[2] == round(27 / 60, 2) and total[3] == 0.2 and total[4] == 0.1

        cl.cookies.clear()
        assert cl.post("/login", data={"code": scode}).status_code == 200
        r = cl.get("/admin/backup/full.zip", follow_redirects=False)
        assert r.status_code != 200 or r.headers.get("content-type") != "application/zip"
