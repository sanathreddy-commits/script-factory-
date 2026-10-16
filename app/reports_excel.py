import io
import time
from datetime import datetime
from zoneinfo import ZoneInfo
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IST = ZoneInfo("Asia/Kolkata")

HEADER_FILL = PatternFill(start_color="00A896", end_color="00A896", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="64748B")
DATA_FONT = Font(name="Segoe UI", size=10, color="1E293B")
BOLD_DATA_FONT = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0")
)


def _style_headers(ws, row_idx, headers):
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[row_idx].height = 28


def _auto_fit_columns(ws, max_cols=20):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or "")
            if "\n" in val:
                val = max(val.split("\n"), key=len)
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(50, max(max_len + 4, 12))


def _embed_proof(ws, cell_ref, proof_name):
    """Put the proof photo INSIDE the cell as a small JPEG thumbnail. Returns True when embedded."""
    try:
        from PIL import Image as PILImage, ImageOps
        from openpyxl.drawing.image import Image as XLImage
        from . import storage
        data = storage.get_proof(proof_name)
        if not data:
            return False
        im = ImageOps.exif_transpose(PILImage.open(io.BytesIO(data))).convert("RGB")
        im.thumbnail((170, 120))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=72)
        buf.seek(0)
        xl = XLImage(buf)
        xl.width, xl.height = im.size
        ws.add_image(xl, cell_ref)
        return True
    except Exception as e:
        print(f"[Excel] proof embed failed for {proof_name}: {e}")
        return False


def generate_master_excel_report(c, lang_id=None) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    gen_time_str = datetime.now(IST).strftime("%d-%b-%Y %I:%M %p IST")

    # =========================================================================
    # SHEET 1: OVERVIEW & SUMMARY
    # =========================================================================
    ws_ov = wb.create_sheet(title="Overview")
    ws_ov.views.sheetView[0].showGridLines = True

    ws_ov.cell(row=1, column=1, value="ScriptFlow - Master Operations Summary").font = TITLE_FONT
    ws_ov.cell(row=2, column=1, value=f"Generated on {gen_time_str}").font = SUBTITLE_FONT

    # Query overview metrics
    langs = c.execute("SELECT * FROM languages ORDER BY id").fetchall()

    tot_rec_min = c.execute(
        "SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE conf_min IS NOT NULL AND status NOT IN ('ASSIGNED','RELEASED','ABANDONED')"
    ).fetchone()["m"]
    tot_rec_h = round(tot_rec_min / 60.0, 1)

    tot_acc_min = c.execute(
        "SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE (review_status IN ('APPROVED','ACCEPTED') OR status IN ('APPROVED','ACCEPTED')) AND conf_min IS NOT NULL"
    ).fetchone()["m"]
    tot_acc_h = round(tot_acc_min / 60.0, 1)

    tot_agents = c.execute("SELECT COUNT(*) n FROM users WHERE role='participant'").fetchone()["n"]
    act_agents = c.execute("SELECT COUNT(*) n FROM users WHERE role='participant' AND active=1").fetchone()["n"]
    tot_pairs = c.execute("SELECT COUNT(*) n FROM pairs WHERE status='ACTIVE'").fetchone()["n"]
    tot_scripts = c.execute("SELECT COUNT(*) n FROM scripts").fetchone()["n"]
    tot_subs = c.execute("SELECT COUNT(*) n FROM assignments WHERE conf_min IS NOT NULL AND status NOT IN ('ASSIGNED','RELEASED','ABANDONED')").fetchone()["n"]
    tot_acc_subs = c.execute("SELECT COUNT(*) n FROM assignments WHERE (review_status IN ('APPROVED','ACCEPTED') OR status IN ('APPROVED','ACCEPTED')) AND conf_min IS NOT NULL").fetchone()["n"]
    tot_pend_subs = c.execute("SELECT COUNT(*) n FROM assignments WHERE status IN ('SUBMITTED','COMPLETED') AND (review_status='PENDING' OR review_status IS NULL) AND conf_min IS NOT NULL").fetchone()["n"]
    tot_rej_subs = c.execute("SELECT COUNT(*) n FROM assignments WHERE (review_status='REJECTED' OR status='REJECTED') AND conf_min IS NOT NULL").fetchone()["n"]

    ov_metrics = [
        ("Total Recorded Hours (Confirmed)", f"{tot_rec_h} h"),
        ("Total Accepted Hours (Quality Verified)", f"{tot_acc_h} h"),
        ("Total Operational Languages", len(langs)),
        ("Total Registered Agents", tot_agents),
        ("Active Agents", act_agents),
        ("Active Formed Pairs", tot_pairs),
        ("Total Scripts in System Library", tot_scripts),
        ("Total Audio Submissions Recorded", tot_subs),
        ("Accepted Submissions", tot_acc_subs),
        ("Submissions Pending QC Review", tot_pend_subs),
        ("Rejected Submissions", tot_rej_subs)
    ]

    _style_headers(ws_ov, 4, ["Metric / Operational Dimension", "Current Value / Stat"])
    for idx, (m, v) in enumerate(ov_metrics, 5):
        cell_m = ws_ov.cell(row=idx, column=1, value=m)
        cell_v = ws_ov.cell(row=idx, column=2, value=v)
        cell_m.font = DATA_FONT
        cell_v.font = BOLD_DATA_FONT
        cell_m.border = THIN_BORDER
        cell_v.border = THIN_BORDER
        cell_m.fill = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL
        cell_v.fill = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL
        cell_v.alignment = Alignment(horizontal="right")
        ws_ov.row_dimensions[idx].height = 22

    _auto_fit_columns(ws_ov)

    # =========================================================================
    # SHEET 2: LANGUAGES
    # =========================================================================
    ws_lang = wb.create_sheet(title="Languages")
    ws_lang.views.sheetView[0].showGridLines = True
    lang_headers = [
        "Sl. No", "Language Code", "Language Name", "Total Scripts",
        "Recorded Hours", "Accepted Hours", "Completion %",
        "Total Agents", "Active Pairs", "Total Scripts", "Available Scripts"
    ]
    _style_headers(ws_lang, 1, lang_headers)

    for idx, l in enumerate(langs, 1):
        lid = l["id"]
        rec_m = c.execute(
            "SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments a JOIN scripts s ON s.id=a.script_id WHERE s.language_id=? AND a.conf_min IS NOT NULL AND a.status NOT IN ('ASSIGNED','RELEASED','ABANDONED')",
            (lid,)
        ).fetchone()["m"]
        acc_m = c.execute(
            "SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments a JOIN scripts s ON s.id=a.script_id WHERE s.language_id=? AND (a.review_status IN ('APPROVED','ACCEPTED') OR a.status IN ('APPROVED','ACCEPTED')) AND a.conf_min IS NOT NULL",
            (lid,)
        ).fetchone()["m"]
        rec_h = round(rec_m / 60.0, 1)
        acc_h = round(acc_m / 60.0, 1)
        tgt = c.execute("SELECT COUNT(*) n FROM scripts WHERE language_id=?", (lid,)).fetchone()["n"]
        done_n = c.execute("SELECT COUNT(DISTINCT a.script_id) n FROM assignments a JOIN scripts s ON s.id=a.script_id WHERE s.language_id=? AND (a.review_status IN ('APPROVED','ACCEPTED') OR a.status IN ('APPROVED','ACCEPTED'))", (lid,)).fetchone()["n"]
        comp_pct = round(done_n / tgt * 100, 1) if tgt else 0

        ag_cnt = c.execute("SELECT COUNT(*) n FROM users WHERE role='participant' AND language_id=?", (lid,)).fetchone()["n"]
        pr_cnt = c.execute("SELECT COUNT(*) n FROM pairs WHERE language_id=? AND status='ACTIVE'", (lid,)).fetchone()["n"]
        sc_cnt = c.execute("SELECT COUNT(*) n FROM scripts WHERE language_id=?", (lid,)).fetchone()["n"]
        av_cnt = c.execute("SELECT COUNT(*) n FROM scripts WHERE language_id=? AND status IN ('READY','AVAILABLE')", (lid,)).fetchone()["n"]

        row_vals = [
            idx, l["code"], l["name"], tgt,
            f"{rec_h} h", f"{acc_h} h", f"{comp_pct}%",
            ag_cnt, pr_cnt, sc_cnt, av_cnt
        ]
        row_num = idx + 1
        fill = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_lang.cell(row=row_num, column=c_idx, value=val)
            cell.font = BOLD_DATA_FONT if c_idx in (2, 3) else DATA_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if c_idx in (1, 2, 4, 5, 6, 7, 8, 9, 10, 11):
                cell.alignment = Alignment(horizontal="center")
        ws_lang.row_dimensions[row_num].height = 20

    _auto_fit_columns(ws_lang)

    # =========================================================================
    # SHEET 3: AGENTS
    # =========================================================================
    ws_ag = wb.create_sheet(title="Agents")
    ws_ag.views.sheetView[0].showGridLines = True
    ag_headers = [
        "Sl. No", "Access Code", "Agent Name", "Phone", "Email",
        "Gender", "Language", "Availability", "Status", "Current Pair / Partner",
        "Assigned Scripts", "Completed Scripts", "Pending Scripts",
        "Recorded Hours", "Cap Limit", "Accepted Hours"
    ]
    _style_headers(ws_ag, 1, ag_headers)

    where_ag = "WHERE u.role='participant'"
    params_ag = []
    if lang_id:
        where_ag += " AND u.language_id=?"
        params_ag.append(lang_id)

    agents = c.execute(
        f"SELECT u.*, l.name as lang_name FROM users u LEFT JOIN languages l ON l.id=u.language_id {where_ag} ORDER BY u.language_id, u.name ASC",
        params_ag
    ).fetchall()

    for idx, ag in enumerate(agents, 1):
        uid = ag["id"]
        # Find partner
        p = c.execute("SELECT * FROM pairs WHERE status='ACTIVE' AND (a=? OR b=?)", (uid, uid)).fetchone()
        partner_name = "Unpaired"
        if p:
            part_id = p["b"] if p["a"] == uid else p["a"]
            p_row = c.execute("SELECT name FROM users WHERE id=?", (part_id,)).fetchone()
            if p_row:
                partner_name = p_row["name"]

        # Scripts and hours
        asg_c = c.execute(
            "SELECT COUNT(*) n FROM assignments WHERE (ua=? OR ub=? OR pair_id IN (SELECT id FROM pairs WHERE a=? OR b=?)) AND status NOT IN ('RELEASED','ABANDONED')",
            (uid, uid, uid, uid)
        ).fetchone()["n"]
        cmp_c = c.execute(
            "SELECT COUNT(*) n FROM assignments WHERE (ua=? OR ub=? OR pair_id IN (SELECT id FROM pairs WHERE a=? OR b=?)) AND (status IN ('SUBMITTED','COMPLETED','CONFIRMED','ACCEPTED') OR review_status IN ('APPROVED','ACCEPTED'))",
            (uid, uid, uid, uid)
        ).fetchone()["n"]
        pnd_c = max(0, asg_c - cmp_c)

        rec_m = c.execute(
            "SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE (ua=? OR ub=? OR pair_id IN (SELECT id FROM pairs WHERE a=? OR b=?)) AND conf_min IS NOT NULL AND status NOT IN ('ASSIGNED','RELEASED','ABANDONED')",
            (uid, uid, uid, uid)
        ).fetchone()["m"]
        acc_m = c.execute(
            "SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE (ua=? OR ub=? OR pair_id IN (SELECT id FROM pairs WHERE a=? OR b=?)) AND (review_status IN ('APPROVED','ACCEPTED') OR status IN ('APPROVED','ACCEPTED')) AND conf_min IS NOT NULL",
            (uid, uid, uid, uid)
        ).fetchone()["m"]

        rec_h = round(rec_m / 60.0, 1)
        acc_h = round(acc_m / 60.0, 1)

        row_vals = [
            idx, ag["code"], ag["name"], ag["phone"] or "-", ag["email"] or "-",
            ag["gender"] or "-", ag["lang_name"] or "-", ag["availability"] or "All day",
            "Active" if ag["active"] else "Inactive", partner_name,
            asg_c, cmp_c, pnd_c,
            f"{rec_h} h", "15.0 h", f"{acc_h} h"
        ]
        row_num = idx + 1
        fill = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_ag.cell(row=row_num, column=c_idx, value=val)
            cell.font = BOLD_DATA_FONT if c_idx in (2, 3) else DATA_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if c_idx in (1, 2, 6, 8, 9, 11, 12, 13, 14, 15, 16):
                cell.alignment = Alignment(horizontal="center")
        ws_ag.row_dimensions[row_num].height = 20

    _auto_fit_columns(ws_ag)

    # =========================================================================
    # SHEET 4: PAIRS
    # =========================================================================
    ws_pr = wb.create_sheet(title="Pairs")
    ws_pr.views.sheetView[0].showGridLines = True
    pr_headers = [
        "Sl. No", "Pair ID", "Language", "Agent A", "Agent A (15h)",
        "Agent B", "Agent B (15h)", "Availability", "Pair Status",
        "Assigned Scripts", "Completed Scripts", "Pending Scripts",
        "Total Recorded Hours", "Pair Cap", "Accepted Hours"
    ]
    _style_headers(ws_pr, 1, pr_headers)

    where_pr = "WHERE p.status='ACTIVE'"
    params_pr = []
    if lang_id:
        where_pr += " AND p.language_id=?"
        params_pr.append(lang_id)

    pairs = c.execute(
        f"SELECT p.*, ua.name as name_a, ua.gender as gender_a, ub.name as name_b, ub.gender as gender_b, l.name as lang_name "
        f"FROM pairs p JOIN users ua ON ua.id=p.a JOIN users ub ON ub.id=p.b JOIN languages l ON l.id=p.language_id "
        f"{where_pr} ORDER BY p.id DESC",
        params_pr
    ).fetchall()

    for idx, p in enumerate(pairs, 1):
        pid = p["id"]
        a_c = c.execute("SELECT COUNT(*) n FROM assignments WHERE pair_id=? AND status NOT IN ('RELEASED','ABANDONED')", (pid,)).fetchone()["n"]
        c_c = c.execute("SELECT COUNT(*) n FROM assignments WHERE pair_id=? AND (status IN ('SUBMITTED','COMPLETED','CONFIRMED','ACCEPTED') OR review_status IN ('APPROVED','ACCEPTED'))", (pid,)).fetchone()["n"]
        p_c = max(0, a_c - c_c)

        tot_m = c.execute("SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE pair_id=? AND conf_min IS NOT NULL AND status NOT IN ('ASSIGNED','RELEASED','ABANDONED')", (pid,)).fetchone()["m"]
        app_m = c.execute("SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE pair_id=? AND (review_status IN ('APPROVED','ACCEPTED') OR status IN ('APPROVED','ACCEPTED')) AND conf_min IS NOT NULL", (pid,)).fetchone()["m"]

        rec_a = c.execute("SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE (ua=? OR ub=?) AND conf_min IS NOT NULL AND status NOT IN ('ASSIGNED','RELEASED','ABANDONED')", (p["a"], p["a"])).fetchone()["m"]
        rec_b = c.execute("SELECT COALESCE(SUM(conf_min), 0.0) m FROM assignments WHERE (ua=? OR ub=?) AND conf_min IS NOT NULL AND status NOT IN ('ASSIGNED','RELEASED','ABANDONED')", (p["b"], p["b"])).fetchone()["m"]

        tot_h = round(tot_m / 60.0, 1)
        app_h = round(app_m / 60.0, 1)
        rec_a_h = round(rec_a / 60.0, 1)
        rec_b_h = round(rec_b / 60.0, 1)

        row_vals = [
            idx, f"Pair #{pid}", p["lang_name"],
            f"{p['name_a']} ({p['gender_a'] or '-'})", f"{rec_a_h} / 15 h",
            f"{p['name_b']} ({p['gender_b'] or '-'})", f"{rec_b_h} / 15 h",
            p["availability"] or "All day", "Limit Reached" if tot_h >= 7.5 else "Active",
            a_c, c_c, p_c,
            f"{tot_h} h", "7.5 h", f"{app_h} h"
        ]
        row_num = idx + 1
        fill = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_pr.cell(row=row_num, column=c_idx, value=val)
            cell.font = BOLD_DATA_FONT if c_idx in (2, 4, 6) else DATA_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if c_idx in (1, 2, 5, 7, 8, 9, 10, 11, 12, 13, 14, 15):
                cell.alignment = Alignment(horizontal="center")
        ws_pr.row_dimensions[row_num].height = 20

    _auto_fit_columns(ws_pr)

    # =========================================================================
    # SHEET 5: SCRIPTS
    # =========================================================================
    ws_sc = wb.create_sheet(title="Scripts")
    ws_sc.views.sheetView[0].showGridLines = True
    sc_headers = [
        "Sl. No", "Script Code", "Language", "Domain", "Subdomain",
        "Specialty", "Word Count", "Status", "Current Assigned Pair"
    ]
    _style_headers(ws_sc, 1, sc_headers)

    where_sc = ""
    params_sc = []
    if lang_id:
        where_sc = "WHERE s.language_id=?"
        params_sc.append(lang_id)

    scripts = c.execute(
        f"SELECT s.*, l.name as lang_name FROM scripts s JOIN languages l ON l.id=s.language_id {where_sc} ORDER BY s.id DESC",
        params_sc
    ).fetchall()

    for idx, sc in enumerate(scripts, 1):
        asg = c.execute(
            "SELECT p.id as pid, ua.name as na, ub.name as nb, a.status as a_stat, a.review_status "
            "FROM assignments a LEFT JOIN pairs p ON p.id=a.pair_id "
            "LEFT JOIN users ua ON ua.id=p.a LEFT JOIN users ub ON ub.id=p.b "
            "WHERE a.script_id=? AND a.status NOT IN ('RELEASED','ABANDONED') ORDER BY a.id DESC LIMIT 1",
            (sc["id"],)
        ).fetchone()

        if asg:
            pair_disp = f"{asg['na']} & {asg['nb']}" if (asg["na"] and asg["nb"]) else f"Pair #{asg['pid']}"
            rev_stat = (asg["review_status"] or "").upper()
            raw_stat = (asg["a_stat"] or "").upper()
            if rev_stat in ("APPROVED", "ACCEPTED") or raw_stat in ("APPROVED", "ACCEPTED", "CONFIRMED"):
                stat_disp = "ACCEPTED"
            elif rev_stat == "REJECTED" or raw_stat == "REJECTED":
                stat_disp = "REJECTED"
            elif raw_stat in ("SUBMITTED", "COMPLETED"):
                stat_disp = "SUBMITTED"
            elif raw_stat in ("ASSIGNED", "IN_SESSION"):
                stat_disp = "ASSIGNED"
            else:
                stat_disp = raw_stat or "ASSIGNED"
        else:
            pair_disp = "-"
            stat_disp = "AVAILABLE" if sc["status"] in ("READY", "AVAILABLE") else sc["status"]

        row_vals = [
            idx, sc["code"], sc["lang_name"], sc["domain"] or "-", sc["subdomain"] or "-",
            sc["specialisation"] or "-", sc["words"] or 0, stat_disp, pair_disp
        ]
        row_num = idx + 1
        fill = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_sc.cell(row=row_num, column=c_idx, value=val)
            cell.font = BOLD_DATA_FONT if c_idx == 2 else DATA_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if c_idx in (1, 2, 7, 8):
                cell.alignment = Alignment(horizontal="center")
        ws_sc.row_dimensions[row_num].height = 20

    _auto_fit_columns(ws_sc)

    # =========================================================================
    # SHEET 6: SUBMISSIONS & QC REPORTS
    # =========================================================================
    ws_rep = wb.create_sheet(title="Submissions & QC")
    ws_rep.views.sheetView[0].showGridLines = True
    rep_headers = [
        "Sl. No", "Assignment ID", "Script Code", "Language", "Speaker A", "Speaker B",
        "Pair ID", "Duration (mm:ss)", "Duration (Mins)", "Review Status",
        "Proof Photo", "QC Review Notes", "Submitted At", "Reviewed At"
    ]
    _style_headers(ws_rep, 1, rep_headers)

    where_rep = "WHERE a.conf_min IS NOT NULL AND a.status NOT IN ('ASSIGNED','RELEASED','ABANDONED')"
    params_rep = []
    if lang_id:
        where_rep += " AND s.language_id=?"
        params_rep.append(lang_id)

    reports = c.execute(
        f"SELECT a.id, a.conf_min, a.proof_img, a.submitted_at, a.updated_at, a.review_status, a.review_note, a.status as a_status, a.pair_id, "
        f"s.code as script_code, ua.name as name_a, ub.name as name_b, l.name as lang_name "
        f"FROM assignments a JOIN scripts s ON s.id=a.script_id JOIN languages l ON l.id=s.language_id "
        f"LEFT JOIN users ua ON ua.id=a.ua LEFT JOIN users ub ON ub.id=a.ub "
        f"{where_rep} ORDER BY a.id DESC",
        params_rep
    ).fetchall()

    for idx, r in enumerate(reports, 1):
        m, s = divmod(int(round(float(r["conf_min"] or 0) * 60)), 60)
        time_str = f"{m:02d}:{s:02d}"

        rev_stat = (r["review_status"] or "").upper()
        raw_stat = (r["a_status"] or "").upper()
        if rev_stat in ("APPROVED", "ACCEPTED") or raw_stat in ("APPROVED", "ACCEPTED"):
            status_disp = "ACCEPTED"
        elif rev_stat == "REJECTED" or raw_stat == "REJECTED":
            status_disp = "REJECTED"
        else:
            status_disp = "SUBMITTED"

        sub_dt = datetime.fromtimestamp(r["submitted_at"], IST).strftime("%Y-%m-%d %H:%M") if r["submitted_at"] else "-"
        rev_dt = datetime.fromtimestamp(r["updated_at"], IST).strftime("%Y-%m-%d %H:%M") if (r["updated_at"] and status_disp in ("ACCEPTED", "REJECTED")) else "-"

        row_vals = [
            idx, r["id"], r["script_code"], r["lang_name"],
            r["name_a"] or "Speaker A", r["name_b"] or "Speaker B",
            f"Pair #{r['pair_id']}", time_str, r["conf_min"],
            status_disp, r["proof_img"] or "None", r["review_note"] or "-",
            sub_dt, rev_dt
        ]
        row_num = idx + 1
        fill = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_rep.cell(row=row_num, column=c_idx, value=val)
            cell.font = BOLD_DATA_FONT if c_idx == 3 else DATA_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if c_idx in (1, 2, 3, 7, 8, 9, 10, 13, 14):
                cell.alignment = Alignment(horizontal="center")
        ws_rep.row_dimensions[row_num].height = 20
        if r["proof_img"]:
            if _embed_proof(ws_rep, f"K{row_num}", r["proof_img"]):
                ws_rep.cell(row=row_num, column=11, value="")
                ws_rep.row_dimensions[row_num].height = 94
            else:
                ws_rep.cell(row=row_num, column=11, value="Photo missing from storage")

    _auto_fit_columns(ws_rep)
    ws_rep.column_dimensions["K"].width = 26

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ======================================================================================
# Clean, professional, FILTERED session export (date range + language).
# Columns mirror the team's manual tracker so a download has everything they need.
# ======================================================================================

def _day_bounds(date_from, date_to):
    """'YYYY-MM-DD' strings -> (start_ts, end_ts) epoch seconds in IST, or (None, None)."""
    start = end = None
    try:
        if date_from:
            d = datetime.strptime(date_from.strip(), "%Y-%m-%d").replace(tzinfo=IST)
            start = d.timestamp()
    except Exception:
        start = None
    try:
        if date_to:
            d = datetime.strptime(date_to.strip(), "%Y-%m-%d").replace(tzinfo=IST)
            end = d.timestamp() + 86400  # inclusive of the whole 'to' day
    except Exception:
        end = None
    return start, end


def _mmss(conf_min):
    try:
        total = int(round(float(conf_min or 0) * 60))
    except (TypeError, ValueError):
        return "-"
    if total <= 0:
        return "-"
    return f"{total // 60:02d}:{total % 60:02d}"


def _fetch_sessions(c, lang_id=None, date_from=None, date_to=None, sup_lang_ids=None):
    where, args = [], []
    if sup_lang_ids:
        where.append("s.language_id IN (%s)" % ",".join("?" * len(sup_lang_ids)))
        args.extend(sup_lang_ids)
    elif lang_id:
        where.append("s.language_id = ?")
        args.append(int(lang_id))
    start, end = _day_bounds(date_from, date_to)
    if start is not None:
        where.append("COALESCE(a.submitted_at, a.assigned_at) >= ?")
        args.append(start)
    if end is not None:
        where.append("COALESCE(a.submitted_at, a.assigned_at) < ?")
        args.append(end)
    where_sql = (" WHERE " + " AND ".join(where)) if where else ""
    sql = (
        "SELECT a.id, s.code, s.domain, s.subdomain, s.specialisation, l.name AS lang, "
        "ua.name AS na, ua.email AS ea, ub.name AS nb, ub.email AS eb, "
        "a.conf_min, a.status, a.review_status, a.review_note, a.proof_img, a.proof_drive_url, a.submitted_at, a.assigned_at "
        "FROM assignments a JOIN scripts s ON s.id=a.script_id JOIN languages l ON l.id=s.language_id "
        "LEFT JOIN users ua ON ua.id=a.ua LEFT JOIN users ub ON ub.id=a.ub "
        + where_sql + " ORDER BY COALESCE(a.submitted_at, a.assigned_at) DESC, a.id DESC"
    )
    return c.execute(sql, tuple(args)).fetchall()


def _who(name, email):
    name = (name or "").strip()
    email = (email or "").strip()
    if name and email:
        return f"{name} <{email}>"
    return name or email or "-"


def _review_label(review_status, status):
    rs = (review_status or "").upper()
    if rs in ("APPROVED", "ACCEPTED"):
        return "Accepted"
    if rs == "REJECTED" or (status or "").upper() == "REJECTED":
        return "Rejected"
    return "Pending"


LINK_FONT = Font(name="Segoe UI", size=10, color="2563EB", underline="single")


def _proof_link(r):
    """Clickable URL for a proof photo: the Google Drive link if present, else a Supabase
    signed URL, else ''. Returns (url, label)."""
    drive_url = (r["proof_drive_url"] or "").strip() if "proof_drive_url" in r.keys() else ""
    if drive_url:
        return drive_url, "View photo"
    name = (r["proof_img"] or "").strip()
    if name:
        try:
            from . import storage
            su = storage.signed_url(name)
            if su:
                return su, "View photo"
        except Exception:
            pass
        return "", name  # no link available, show filename
    return "", "-"


def generate_sessions_report(c, lang_id=None, date_from=None, date_to=None, sup_lang_ids=None) -> bytes:
    """A clean single-focus workbook: a professional 'Sessions' sheet (team-tracker layout)
    plus a compact 'Summary' sheet, both honouring the language + date-range filters."""
    rows = _fetch_sessions(c, lang_id, date_from, date_to, sup_lang_ids)
    wb = openpyxl.Workbook()

    # ---- Sessions sheet ----
    ws = wb.active
    ws.title = "Sessions"
    scope = "All languages"
    if sup_lang_ids:
        names = [r["name"] for r in c.execute(
            "SELECT name FROM languages WHERE id IN (%s) ORDER BY name" % ",".join("?" * len(sup_lang_ids)),
            tuple(sup_lang_ids)).fetchall()]
        scope = ", ".join(names) if names else "Assigned languages"
    elif lang_id:
        lr = c.execute("SELECT name FROM languages WHERE id=?", (int(lang_id),)).fetchone()
        scope = lr["name"] if lr else f"Language #{lang_id}"
    rng = "All dates"
    if date_from or date_to:
        rng = f"{date_from or 'start'} to {date_to or 'today'}"

    ws.cell(row=1, column=1, value="Session / Recording Report").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Language: {scope}   |   Period: {rng}   |   "
            f"Generated: {datetime.now(IST).strftime('%d-%b-%Y %I:%M %p IST')}   |   Rows: {len(rows)}").font = SUBTITLE_FONT

    headers = ["S.No", "Date", "Language", "Topic", "Script Code", "User 1 (Host)",
               "User 2 (Guest)", "Duration (mm:ss)", "Status", "Review", "Comment / Invalid Reason", "Proof"]
    hdr_row = 4
    _style_headers(ws, hdr_row, headers)
    for i, r in enumerate(rows):
        rr = hdr_row + 1 + i
        ts = r["submitted_at"] or r["assigned_at"]
        date_str = datetime.fromtimestamp(ts, IST).strftime("%d-%b-%Y %H:%M") if ts else "-"
        topic = " / ".join([x for x in [r["domain"], r["subdomain"], r["specialisation"]] if x])
        proof_url, proof_label = _proof_link(r)
        vals = [i + 1, date_str, r["lang"], topic or "-", r["code"], _who(r["na"], r["ea"]),
                _who(r["nb"], r["eb"]), _mmss(r["conf_min"]), (r["status"] or "-"),
                _review_label(r["review_status"], r["status"]), (r["review_note"] or "").strip() or "-",
                proof_label]
        fill = ZEBRA_FILL if i % 2 == 0 else WHITE_FILL
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=rr, column=ci, value=v)
            cell.font = DATA_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if ci in (1, 2, 8, 9, 10, 12) else "left",
                                       vertical="top", wrap_text=ci in (4, 11))
            if ci == 12 and proof_url:
                cell.hyperlink = proof_url
                cell.font = LINK_FONT
    if not rows:
        ws.cell(row=hdr_row + 1, column=1, value="No sessions match this filter.").font = DATA_FONT
    _auto_fit_columns(ws)
    ws.column_dimensions["D"].width = 26
    ws.column_dimensions["K"].width = 42
    ws.freeze_panes = f"A{hdr_row + 1}"
    ws.auto_filter.ref = f"A{hdr_row}:L{hdr_row + max(len(rows), 1)}"

    # ---- Summary sheet (per language, within the same filter) ----
    ws2 = wb.create_sheet("Summary")
    ws2.cell(row=1, column=1, value="Summary").font = TITLE_FONT
    ws2.cell(row=2, column=1, value=f"Language: {scope}   |   Period: {rng}").font = SUBTITLE_FONT
    s_hdr = ["Language", "Sessions", "Accepted", "Pending", "Rejected", "Total minutes"]
    _style_headers(ws2, 4, s_hdr)
    agg = {}
    for r in rows:
        a = agg.setdefault(r["lang"], {"n": 0, "acc": 0, "pen": 0, "rej": 0, "min": 0.0})
        a["n"] += 1
        lab = _review_label(r["review_status"], r["status"])
        a["acc" if lab == "Accepted" else "rej" if lab == "Rejected" else "pen"] += 1
        try:
            a["min"] += float(r["conf_min"] or 0)
        except (TypeError, ValueError):
            pass
    for i, (lang, a) in enumerate(sorted(agg.items())):
        rr = 5 + i
        for ci, v in enumerate([lang, a["n"], a["acc"], a["pen"], a["rej"], round(a["min"], 1)], 1):
            cell = ws2.cell(row=rr, column=ci, value=v)
            cell.font = DATA_FONT
            cell.fill = ZEBRA_FILL if i % 2 == 0 else WHITE_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center" if ci > 1 else "left")
    _auto_fit_columns(ws2)
    ws2.freeze_panes = "A5"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
