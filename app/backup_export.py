"""Full backup download: one ZIP with a readable Excel report and a raw CSV of every important table.

One SELECT per table (no per-row queries, no photo downloads), so a backup costs about the size of the data in
database egress. Secrets never leave: passwords, login codes, API keys, connection configs and settings are skipped,
and the ~6 KB duplicate-detection fingerprint (scripts.shingles) is left out.
"""
import csv
import io
import zipfile
from datetime import datetime
from zoneinfo import ZoneInfo

from . import db

IST = ZoneInfo("Asia/Kolkata")

# table -> columns to export (None = all columns). Anything not listed here is not exported.
TABLES = {
    "languages": None,
    "users": ["id", "name", "phone", "email", "role", "language_id", "active", "gender", "perm", "consent_at", "consent_v",
              "optin", "availability", "created_at", "last_seen"],
    "supervisor_languages": None,
    "pairs": None,
    "scripts": [c.split(".", 1)[1] for c in db.SCRIPT_COLS.split(", ")],
    "script_versions": None,
    "assignments": None,
    "ledger": None,
    "reviews": None,
    "issues": None,
    "quotas": None,
    "audit": None,
}


def _ist(ts):
    try:
        return datetime.fromtimestamp(float(ts), IST).strftime("%d-%b-%Y %H:%M") if ts else ""
    except (TypeError, ValueError, OSError):
        return ""


def _rows(c, table, cols):
    sel = ", ".join(cols) if cols else "*"
    cur = c.execute(f"SELECT {sel} FROM {table} ORDER BY 1")
    rows = cur.fetchall()
    if rows:
        return list(rows[0].keys()), [tuple(r) for r in rows]
    return list(cols or []), []


def _csv(header, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def _review(review_status, status):
    rs, st = (review_status or "").upper(), (status or "").upper()
    if rs in ("APPROVED", "ACCEPTED") or st in ("APPROVED", "ACCEPTED", "CONFIRMED"):
        return "Accepted"
    if rs == "REJECTED" or st in ("REJECTED", "QC_FAIL"):
        return "Rejected"
    return "Pending" if st in ("SUBMITTED", "COMPLETED") else ""


def _excel(data):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B4D3E")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        for i, h in enumerate(header):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = (widths or {}).get(h, max(10, min(40, len(str(h)) + 4)))
        return ws

    langs = {r[0]: r for r in data["languages"][1]}
    lang_name = {k: v[data["languages"][0].index("name")] for k, v in langs.items()}
    uh, urows = data["users"]
    users = {r[0]: dict(zip(uh, r)) for r in urows}
    sh, srows = data["scripts"]
    scripts = {r[0]: dict(zip(sh, r)) for r in srows}
    ah, arows = data["assignments"]
    asg = [dict(zip(ah, r)) for r in arows]

    # Work done: one row per recording that was submitted
    work = []
    per_lang = {}
    for a in asg:
        if not a.get("conf_min") or (a.get("status") or "") in ("ASSIGNED", "RELEASED", "ABANDONED", "IN_SESSION"):
            continue
        s = scripts.get(a["script_id"], {})
        lid = s.get("language_id")
        rv = _review(a.get("review_status"), a.get("status"))
        ua, ub = users.get(a.get("ua"), {}), users.get(a.get("ub"), {})
        mins = float(a["conf_min"] or 0)
        work.append((_ist(a.get("submitted_at") or a.get("assigned_at")), lang_name.get(lid, ""), s.get("code", ""),
                     s.get("domain", ""), s.get("subdomain", ""), a.get("pair_id"), ua.get("name", ""), ub.get("name", ""),
                     round(mins, 2), a.get("status", ""), rv, (a.get("review_note") or "").strip(),
                     (a.get("proof_drive_url") or "").strip(), (a.get("proof_img") or "").strip()))
        p = per_lang.setdefault(lid, {"n": 0, "rec": 0.0, "acc": 0.0, "rej": 0.0, "pend": 0.0})
        p["n"] += 1
        p["rec"] += mins / 60
        p["acc" if rv == "Accepted" else "rej" if rv == "Rejected" else "pend"] += mins / 60

    ph, prows = data["pairs"]
    pairs_by_lang = {}
    for r in prows:
        d = dict(zip(ph, r))
        pairs_by_lang.setdefault(d["language_id"], []).append(d)

    summary = [("Generated", datetime.now(IST).strftime("%d-%b-%Y %I:%M %p IST"), "", "", "", "", "", ""), ("", "", "", "", "", "", "", ""),
               ("Language", "Recordings", "Recorded hrs", "Accepted hrs", "Rejected hrs", "Pending review hrs", "Active pairs", "Scripts")]
    scripts_by_lang = {}
    for s in scripts.values():
        scripts_by_lang[s["language_id"]] = scripts_by_lang.get(s["language_id"], 0) + 1
    tot = [0, 0.0, 0.0, 0.0, 0.0, 0, 0]
    for lid in sorted(lang_name, key=lambda k: lang_name[k] or ""):
        p = per_lang.get(lid, {"n": 0, "rec": 0.0, "acc": 0.0, "rej": 0.0, "pend": 0.0})
        act = sum(1 for x in pairs_by_lang.get(lid, []) if (x.get("status") or "") == "ACTIVE")
        nsc = scripts_by_lang.get(lid, 0)
        if not (p["n"] or act or nsc):
            continue
        row = (lang_name[lid], p["n"], round(p["rec"], 2), round(p["acc"], 2), round(p["rej"], 2), round(p["pend"], 2), act, nsc)
        summary.append(row)
        for i, v in enumerate(row[1:]):
            tot[i] += v
    summary.append(("All languages", tot[0], round(tot[1], 2), round(tot[2], 2), round(tot[3], 2), round(tot[4], 2), tot[5], tot[6]))
    ws = wb.create_sheet("Summary")
    for r in summary:
        ws.append(list(r))
    for cell in ws[3]:
        cell.font, cell.fill = head_font, head_fill
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16

    sheet("Work done", ["Submitted (IST)", "Language", "Script", "Domain", "Subdomain", "Pair", "Agent A", "Agent B", "Minutes",
                        "Status", "Review", "Review note", "Drive photo link", "Photo file"], work,
          {"Submitted (IST)": 18, "Script": 18, "Agent A": 22, "Agent B": 22, "Review note": 40, "Drive photo link": 40, "Photo file": 30})
    sheet("Agents", ["ID", "Name", "Phone", "Email", "Role", "Language", "Active", "Gender", "Availability", "Created (IST)", "Last seen (IST)"],
          [(u["id"], u["name"], u["phone"], u["email"], u["role"], lang_name.get(u["language_id"], ""), u["active"], u["gender"],
            u["availability"], _ist(u["created_at"]), _ist(u["last_seen"])) for u in users.values()],
          {"Name": 24, "Email": 28, "Phone": 16})
    sheet("Pairs", ["ID", "Language", "Agent A", "Agent B", "Status", "Availability", "Created (IST)", "Dropped (IST)", "Note"],
          [(d["id"], lang_name.get(d["language_id"], ""), users.get(d["a"], {}).get("name", ""), users.get(d["b"], {}).get("name", ""),
            d["status"], d.get("availability", ""), _ist(d["created_at"]), _ist(d.get("dropped_at")), d.get("note", ""))
           for ds in pairs_by_lang.values() for d in ds], {"Agent A": 22, "Agent B": 22, "Note": 30})
    sheet("Scripts", ["ID", "Language", "Code", "Domain", "Subdomain", "Specialisation", "Status", "Words", "Source", "Created (IST)"],
          [(s["id"], lang_name.get(s["language_id"], ""), s["code"], s["domain"], s["subdomain"], s["specialisation"], s["status"],
            s["words"], s["source"], _ist(s["created_at"])) for s in scripts.values()], {"Code": 18, "Subdomain": 22, "Specialisation": 22})
    sheet("Languages", data["languages"][0], data["languages"][1])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_backup_zip(c):
    """(zip_bytes, file_name)."""
    data = {t: _rows(c, t, cols) for t, cols in TABLES.items()}
    stamp = datetime.now(IST).strftime("%Y-%m-%d_%H%M")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("ScriptFlow_report.xlsx", _excel(data))
        for t, (header, rows) in data.items():
            z.writestr(f"tables/{t}.csv", _csv(header, rows))
        z.writestr("README.txt", (
            f"ScriptFlow full backup, {stamp} IST\n\n"
            "ScriptFlow_report.xlsx  readable report: Summary, Work done (one row per recording), Agents, Pairs, Scripts.\n"
            "tables/*.csv            raw copy of each table (times are Unix seconds; script text is in script_versions.csv).\n\n"
            "Not included on purpose: passwords, login codes, API keys, connection settings, proof photo files\n"
            "(photos stay in Supabase Storage / Google Drive; the report lists each photo's file name and Drive link).\n"))
    return buf.getvalue(), f"ScriptFlow_backup_{stamp}.zip"
